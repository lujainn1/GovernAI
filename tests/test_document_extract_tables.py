"""Tables in DOCX/PPTX uploads (data inventories, risk registers) must reach
the agents - they used to be silently dropped by the extractor."""
from io import BytesIO

from docx import Document
from pptx import Presentation
from pptx.util import Inches

from app.tools.document_extract import extract_text


def _docx_with_table() -> bytes:
    doc = Document()
    doc.add_paragraph("Data Protection Impact Assessment")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Field"
    table.cell(0, 1).text = "Sent to model?"
    table.cell(1, 0).text = "National ID"
    table.cell(1, 1).text = "Yes"
    doc.add_paragraph("Sign-off pending")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_extraction_includes_tables_in_document_order():
    text = extract_text("dpia.docx", _docx_with_table()).text
    lines = text.splitlines()
    assert lines[0] == "Data Protection Impact Assessment"
    assert "National ID | Yes" in lines
    assert lines[-1] == "Sign-off pending"


def test_pptx_extraction_includes_tables():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    shape = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1))
    shape.table.cell(1, 0).text = "R1"
    shape.table.cell(1, 1).text = "Automated rejection"
    buf = BytesIO()
    prs.save(buf)
    text = extract_text("deck.pptx", buf.getvalue()).text
    assert "R1 | Automated rejection" in text
